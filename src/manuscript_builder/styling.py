"""
Styling module for IEEE Black & White Manuscript Generation.
Configures page layout, typography, borders, table shading, and formal boxes.
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def setup_document():
    doc = docx.Document()
    
    # 0.75 in Margins for IEEE standard
    for s in doc.sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        s.page_width = Inches(8.5)
        s.page_height = Inches(11.0)
        
        # Header / Footer
        header = s.header
        h_p = header.paragraphs[0]
        h_p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
        h_run = h_p.add_run("IEEE TRANSACTIONS ON INFORMATION FORENSICS AND SECURITY (DRAFT / MONOGRAPH)")
        h_run.font.name = "Times New Roman"
        h_run.font.size = Pt(8.5)
        h_run.font.color.rgb = RGBColor(120, 120, 120)
        
        footer = s.footer
        f_p = footer.paragraphs[0]
        f_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        f_run = f_p.add_run("PromptGuard AI: Evidential Cascaded Routing & Conformal Risk Control")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(8.5)
        f_run.font.color.rgb = RGBColor(120, 120, 120)
        
    return doc

def set_cell_shading(cell, color_hex="F2F2F2"):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="10" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="10" w:space="0" w:color="000000"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="D3D3D3"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def add_p(doc, text, bold_prefix="", italic=False, space_after=4, align=WD_PARAGRAPH_ALIGNMENT.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.italic = italic
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h0(doc, title_text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title_text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(18)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_part_header(doc, part_title):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.paragraph_format.space_before = Pt(20)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(part_title)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h1(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.italic = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_h3(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_callout(doc, text, title=""):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F9F9F9")
    set_cell_margins(cell, 120, 120, 160, 160)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="000000"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.1
    if title:
        r_t = p.add_run(title + "\n")
        r_t.font.name = "Times New Roman"
        r_t.font.size = Pt(9.5)
        r_t.font.bold = True
        r_t.font.color.rgb = RGBColor(0, 0, 0)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(20, 20, 20)

def add_table_data(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    
    # Header Row
    hdr_row = table.rows[0]
    tblHeader = parse_xml(f'<w:tblHeader {nsdecls("w")}/>')
    hdr_row._tr.get_or_add_trPr().append(tblHeader)
    
    for i, h in enumerate(headers):
        c = hdr_row.cells[i]
        set_cell_shading(c, "EAEAEA")
        set_cell_margins(c, 90, 90, 110, 110)
        p = c.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        
    # Data Rows
    for r_idx, r_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        shading_color = "F9F9F9" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(r_data):
            c = row.cells[c_idx]
            set_cell_shading(c, shading_color)
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT if c_idx == 0 and len(str(val)) > 10 else WD_PARAGRAPH_ALIGNMENT.CENTER
            r = p.add_run(str(val))
            r.font.name = "Times New Roman"
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0, 0, 0)
            
    if col_widths and len(col_widths) == len(headers):
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Inches(w)
                
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_algorithm_box(doc, algo_num, title, inputs, outputs, steps):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "FFFFFF")
    set_cell_margins(cell, 100, 100, 140, 140)
    
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_head = p.add_run(f"Algorithm {algo_num}: {title}\n")
    r_head.font.name = "Times New Roman"
    r_head.font.size = Pt(9.5)
    r_head.font.bold = True
    
    r_io = p.add_run(f"Input: {inputs}\nOutput: {outputs}\n")
    r_io.font.name = "Times New Roman"
    r_io.font.size = Pt(8.5)
    r_io.font.italic = True
    
    for idx, step in enumerate(steps):
        r_step = p.add_run(f"{idx+1:2d}: {step}\n")
        r_step.font.name = "Courier New"
        r_step.font.size = Pt(8)
        
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
